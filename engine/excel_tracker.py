import os
import re
import hashlib
import datetime
from typing import Optional, List, Dict, Any
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .date_utils import get_app_root, clean_filename

class ExcelTracker:
    """
    Enterprise Excel-based Master State Tracker.
    Replaces SQLite with a visual, auditable Microsoft Excel workbook (data/addendum_master_tracker.xlsx).
    Provides in-memory O(1) deduplication by URL, filename, and content MD5 hash.
    """
    
    HEADERS_ADDENDUMS = [
        "Sl No.",
        "AMC Name",
        "Document Title",
        "Publication Date",
        "Original File Name",
        "Original Download Link",
        "Local File Path",
        "File Size (KB)",
        "Download Date",
        "Download Timestamp",
        "File Hash (MD5)",
        "Status"
    ]

    HEADERS_HISTORY = [
        "Run Date",
        "Total AMCs Processed",
        "New Downloads",
        "Already Up to Date",
        "Errors / Exceptions",
        "Execution Timestamp"
    ]

    def __init__(self, tracker_path: str = None):
        if tracker_path is None:
            base_dir = get_app_root()
            tracker_path = os.path.join(base_dir, "data", "addendum_master_tracker.xlsx")
        self.tracker_path = tracker_path
        os.makedirs(os.path.dirname(self.tracker_path), exist_ok=True)

        # In-memory deduplication indices
        self.downloaded_urls = set()
        self.downloaded_files = set()  # (amc_lower, filename_lower)
        self.downloaded_hashes = set()

        self._init_workbook()
        self._load_indices()

    def _get_styles(self):
        header_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='D9D9D9'),
            right=Side(style='thin', color='D9D9D9'),
            top=Side(style='thin', color='D9D9D9'),
            bottom=Side(style='thin', color='D9D9D9')
        )
        alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        link_font = Font(name="Calibri", size=10, color="0563C1", underline="single")
        default_font = Font(name="Calibri", size=10)
        return header_fill, header_font, header_align, thin_border, alt_fill, link_font, default_font

    def _init_workbook(self):
        if not os.path.exists(self.tracker_path):
            wb = Workbook()
            # Setup Sheet 1: Downloaded Addendums
            ws1 = wb.active
            ws1.title = "Downloaded Addendums"
            ws1.append(self.HEADERS_ADDENDUMS)
            ws1.row_dimensions[1].height = 28

            # Setup Sheet 2: Execution History
            ws2 = wb.create_sheet(title="Execution History")
            ws2.append(self.HEADERS_HISTORY)
            ws2.row_dimensions[1].height = 28

            header_fill, header_font, header_align, _, _, _, _ = self._get_styles()
            for col_idx in range(1, len(self.HEADERS_ADDENDUMS) + 1):
                cell = ws1.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            for col_idx in range(1, len(self.HEADERS_HISTORY) + 1):
                cell = ws2.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align

            wb.save(self.tracker_path)

    def _load_indices(self):
        """Loads URLs, filenames, and hashes into memory for ultra-fast deduplication."""
        self.downloaded_urls.clear()
        self.downloaded_files.clear()
        self.downloaded_hashes.clear()

        if not os.path.exists(self.tracker_path):
            return

        try:
            wb = load_workbook(self.tracker_path, read_only=True, data_only=True)
            if "Downloaded Addendums" in wb.sheetnames:
                ws = wb["Downloaded Addendums"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 5:
                        continue
                    amc_name = str(row[1] or "").strip().lower()
                    orig_fname = str(row[4] or "").strip().lower()
                    url = str(row[5] or "").strip()
                    fhash = str(row[10] or "").strip().lower() if len(row) > 10 else ""

                    if url:
                        self.downloaded_urls.add(url)
                    if amc_name and orig_fname:
                        self.downloaded_files.add((amc_name, orig_fname))
                    if fhash:
                        self.downloaded_hashes.add(fhash)
            wb.close()
        except Exception as e:
            print(f"[!] Warning: Could not load index from Excel Tracker: {e}")

    def is_downloaded(self, pdf_url: str, filename: str = None, amc_name: str = None) -> bool:
        """
        Deduplication check:
        1. Checks if the PDF download URL is already recorded.
        2. Checks if the exact filename for that AMC has already been saved.
        """
        if pdf_url and pdf_url.strip() in self.downloaded_urls:
            return True
        if filename and amc_name:
            key = (amc_name.strip().lower(), filename.strip().lower())
            if key in self.downloaded_files:
                return True
        return False

    def is_hash_downloaded(self, file_hash: str) -> bool:
        """Checks if a file with identical binary MD5 hash exists."""
        if not file_hash:
            return False
        return file_hash.strip().lower() in self.downloaded_hashes

    def record_download(
        self,
        amc_id: str,
        amc_name: str,
        doc_title: str,
        doc_date: str,
        pdf_url: str,
        file_hash: str,
        local_filename: str,
        local_path: str,
        file_size_kb: float,
        download_date: str,
        status: str = "Downloaded"
    ) -> bool:
        """Appends a new download row to the Excel Master Tracker."""
        try:
            wb = load_workbook(self.tracker_path)
            if "Downloaded Addendums" not in wb.sheetnames:
                ws = wb.create_sheet(title="Downloaded Addendums")
                ws.append(self.HEADERS_ADDENDUMS)
            else:
                ws = wb["Downloaded Addendums"]

            row_idx = ws.max_row + 1
            ws.row_dimensions[row_idx].height = 20
            sl_no = row_idx - 1

            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            row_values = [
                sl_no,
                amc_name,
                doc_title or "Notice cum Addendum",
                doc_date or "",
                local_filename,
                pdf_url or "",
                local_path or "",
                file_size_kb if file_size_kb else 0.0,
                download_date,
                now_str,
                file_hash or "",
                status
            ]
            ws.append(row_values)

            _, _, _, thin_border, alt_fill, link_font, default_font = self._get_styles()
            is_alt = (sl_no % 2 == 0)

            for col_idx in range(1, len(self.HEADERS_ADDENDUMS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.font = default_font
                if is_alt:
                    cell.fill = alt_fill

                # Hyperlink for URL (Column 6)
                if col_idx == 6 and pdf_url and pdf_url.startswith("http"):
                    cell.font = link_font
                    cell.hyperlink = pdf_url

                # Hyperlink for Local File Path (Column 7)
                if col_idx == 7 and local_path and os.path.exists(local_path):
                    cell.font = link_font
                    cell.hyperlink = local_path

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    v = str(cell.value or "")
                    if len(v) > max_len:
                        max_len = len(v)
                ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 65)

            wb.save(self.tracker_path)
            wb.close()

            # Update in-memory sets
            if pdf_url:
                self.downloaded_urls.add(pdf_url.strip())
            if amc_name and local_filename:
                self.downloaded_files.add((amc_name.strip().lower(), local_filename.strip().lower()))
            if file_hash:
                self.downloaded_hashes.add(file_hash.strip().lower())

            return True
        except Exception as e:
            print(f"[!] Error recording download to Excel tracker: {e}")
            return False

    def get_downloads_for_date(self, download_date: str) -> List[Dict[str, Any]]:
        """Retrieves all downloads matching download_date (YYYY-MM-DD)."""
        records = []
        if not os.path.exists(self.tracker_path):
            return records

        try:
            wb = load_workbook(self.tracker_path, read_only=True, data_only=True)
            if "Downloaded Addendums" in wb.sheetnames:
                ws = wb["Downloaded Addendums"]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or len(row) < 9:
                        continue
                    row_date = str(row[8] or "").strip()
                    if row_date == download_date:
                        records.append({
                            "amc_name": row[1],
                            "doc_title": row[2],
                            "doc_date": row[3],
                            "local_filename": row[4],
                            "pdf_url": row[5],
                            "local_path": row[6],
                            "file_size_kb": row[7],
                            "created_at": row[9],
                            "status": row[11] if len(row) > 11 else "Downloaded"
                        })
            wb.close()
        except Exception as e:
            print(f"[!] Error reading downloads for date from Excel: {e}")

        return records

    def record_run(
        self,
        run_date: str,
        total_amcs: int,
        new_downloads: int,
        skipped: int,
        errors: int,
        summary_json: str = ""
    ):
        """Appends run statistics to the Execution History worksheet."""
        try:
            wb = load_workbook(self.tracker_path)
            if "Execution History" not in wb.sheetnames:
                ws = wb.create_sheet(title="Execution History")
                ws.append(self.HEADERS_HISTORY)
            else:
                ws = wb["Execution History"]

            row_idx = ws.max_row + 1
            now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

            ws.append([
                run_date,
                total_amcs,
                new_downloads,
                skipped,
                errors,
                now_str
            ])

            _, _, _, thin_border, alt_fill, _, default_font = self._get_styles()
            is_alt = ((row_idx - 1) % 2 == 0)

            for col_idx in range(1, len(self.HEADERS_HISTORY) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.font = default_font
                if is_alt:
                    cell.fill = alt_fill

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    v = str(cell.value or "")
                    if len(v) > max_len:
                        max_len = len(v)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

            wb.save(self.tracker_path)
            wb.close()
        except Exception as e:
            print(f"[!] Error recording run history to Excel tracker: {e}")

    def import_local_file(
        self,
        file_path: str,
        amc_name: str = "Unknown AMC",
        publish_date: str = "",
        doc_title: str = ""
    ) -> bool:
        """
        Registers an existing, previously downloaded PDF file into the tracker.
        Computes its MD5 hash and file size, preventing duplicate future downloads.
        """
        if not os.path.exists(file_path):
            return False

        filename = os.path.basename(file_path)
        if not filename.lower().endswith(".pdf"):
            return False

        # Check if already indexed
        if (amc_name.strip().lower(), filename.strip().lower()) in self.downloaded_files:
            return False

        try:
            with open(file_path, "rb") as f:
                content = f.read()

            file_hash = hashlib.md5(content).hexdigest()
            if file_hash in self.downloaded_hashes:
                return False

            file_size_kb = round(len(content) / 1024.0, 2)
            today_str = datetime.date.today().strftime("%Y-%m-%d")

            title = doc_title or os.path.splitext(filename)[0]

            return self.record_download(
                amc_id="imported",
                amc_name=amc_name,
                doc_title=title,
                doc_date=publish_date,
                pdf_url="",
                file_hash=file_hash,
                local_filename=filename,
                local_path=os.path.abspath(file_path),
                file_size_kb=file_size_kb,
                download_date=today_str,
                status="Historical Import"
            )
        except Exception as e:
            print(f"[!] Failed to import {file_path}: {e}")
            return False
